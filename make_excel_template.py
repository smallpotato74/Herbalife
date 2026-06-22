#!/usr/bin/env python3
"""
make_excel_template.py
======================
生成一個期權篩選 Excel 模板:options_screener_template.xlsx

三張表:
  1) 說明        — 點用 + 評分規則
  2) Watchlist   — 每隻標的算 IVR / IV-HV / 方向建議
  3) OptionScorer— 每個候選合約自動計 expected move + 4 項子分 + 總分

所有公式都用 Excel / Google Sheets 通用函數(IF/MIN/MAX/SQRT/LN/AND...),
你可以直接整份貼上 Google Sheets 用。

跑法: python make_excel_template.py
需要: pip install openpyxl
"""

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

HDR_FILL = PatternFill("solid", fgColor="2A75BB")
HDR_FONT = Font(bold=True, color="FFFFFF", size=10)
SUB_FILL = PatternFill("solid", fgColor="DCE6F1")
IN_FILL = PatternFill("solid", fgColor="FFF2CC")   # 黃 = 你要填
OUT_FILL = PatternFill("solid", fgColor="E2EFDA")  # 綠 = 自動計
THIN = Side(style="thin", color="BBBBBB")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


def style_header(ws, row, ncols, start=1):
    for c in range(start, start + ncols):
        cell = ws.cell(row=row, column=c)
        cell.fill = HDR_FILL
        cell.font = HDR_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center",
                                   wrap_text=True)
        cell.border = BORDER


def build():
    wb = Workbook()

    # ===================== 1) 說明 =====================
    ws = wb.active
    ws.title = "說明"
    ws.column_dimensions["A"].width = 100
    lines = [
        ("期權篩選模板  —  用 ATR + IV + Greeks 揾高質期權", True),
        ("⚠️ 教育用途,非投資建議。期權高風險,落場前用模擬倉。", False),
        ("", False),
        ("【顏色】黃色格 = 你要自己填數;  綠色格 = 公式自動計,唔好改。", True),
        ("", False),
        ("【步驟】", True),
        ("1. Watchlist 表:填每隻股嘅 現價/ATR/HV20/當前IV/一年IV高低/earnings日。", False),
        ("   → 自動出 IVR、IV/HV、方向建議(BUY=買期權 / SELL=賣期權收租)。", False),
        ("2. OptionScorer 表:每行填一個候選合約(strike/DTE/delta/OI/量/買賣價/IV)。", False),
        ("   → 自動計 expected move + 4 項子分 + 總分。總分由高到低排就係你嘅清單。", False),
        ("", False),
        ("【數字邊度嚟】", True),
        ("• 現價/ATR/HV:券商或 yfinance(配 options_screener.py 一齊用)。", False),
        ("• 當前IV / 一年IV高低:期權平台(tastytrade/IBKR/barchart 有 IV Rank)。", False),
        ("• Delta/OI/量/IV(逐個合約):券商期權鏈。", False),
        ("", False),
        ("【評分規則(每項 25 分,滿分 100)】", True),
        ("• 流動性:OI 越大、成交越多、買賣價差越窄 → 分越高。", False),
        ("• Delta:落喺目標區間(買 0.55-0.70 / 賣 0.15-0.30)→ 滿分。", False),
        ("• 到價配合:買方 strike 喺 1個標準差內最好;賣方 strike 喺 1SD 外最安全。", False),
        ("• IV regime:買方想 IV/HV 低(平);賣方想 IV/HV 高(貴)。", False),
        ("", False),
        ("【方向判斷】IVR<30 或 IV/HV<1.0 → BUY;  IVR>50 或 IV/HV>1.2 → SELL。", True),
        ("", False),
        ("【退場(寫喺報告,模板唔自動做)】", True),
        ("• 買方:+50~100% 止賺 / -40~50% 止蝕 / 21 DTE 重評。", False),
        ("• 賣方:收 50% premium 止賺 / 蝕 2× credit 止蝕 / 21 DTE 管理避 gamma。", False),
    ]
    for i, (txt, bold) in enumerate(lines, 1):
        c = ws.cell(row=i, column=1, value=txt)
        c.font = Font(bold=bold, size=12 if (bold and i == 1) else 10,
                      color="2A75BB" if bold else "000000")
        c.alignment = Alignment(wrap_text=True, vertical="top")

    # ===================== 2) Watchlist =====================
    ws = wb.create_sheet("Watchlist")
    headers = ["Ticker", "現價", "ATR(14)", "當前IV", "一年IV低", "一年IV高",
               "HV20", "Earnings日", "↓自動↓ IVR", "IV/HV",
               "方向建議", "1SD@30d", "1SD@45d"]
    for j, h in enumerate(headers, 1):
        ws.cell(row=1, column=j, value=h)
    style_header(ws, 1, len(headers))

    # 註: 輸入欄 A-H, 自動欄 I-M
    notes = ["代號", "$", "$", "小數0.30", "小數", "小數",
             "小數0.25", "YYYY-MM-DD", "(IV-低)/(高-低)", "當前IV/HV",
             "BUY/SELL", "S×IV×√(30/365)", "S×IV×√(45/365)"]
    for j, n in enumerate(notes, 1):
        cc = ws.cell(row=2, column=j, value=n)
        cc.font = Font(italic=True, size=8, color="888888")
        cc.alignment = Alignment(horizontal="center")

    # 範例數據(對應 options_screener.py 嘅 demo 設定,方便對照)
    samples = [
        ["AAPL", 195.0, 4.39, 0.22, 0.18, 0.45, 0.234, "2026-07-10"],
        ["NVDA", 120.0, 6.65, 0.72, 0.40, 0.95, 0.544, "2026-07-01"],
        ["KO",   62.0, 0.90, 0.16, 0.12, 0.28, 0.143, ""],
        ["SPY",  540.0, 6.50, 0.12, 0.09, 0.35, 0.158, ""],
    ]
    first_data = 3
    for r, row in enumerate(samples, first_data):
        for j, v in enumerate(row, 1):
            cell = ws.cell(row=r, column=j, value=v)
            cell.fill = IN_FILL
            cell.border = BORDER
        # I: IVR
        ws.cell(row=r, column=9,
                value=f"=IF((F{r}-E{r})=0,\"\",ROUND((D{r}-E{r})/(F{r}-E{r})*100,0))")
        # J: IV/HV
        ws.cell(row=r, column=10, value=f"=IF(G{r}=0,\"\",ROUND(D{r}/G{r},2))")
        # K: 方向建議
        ws.cell(row=r, column=11,
                value=(f"=IF(OR(I{r}>50,J{r}>1.2),\"SELL\","
                       f"IF(OR(I{r}<30,J{r}<1),\"BUY\",\"NEUTRAL\"))"))
        # L: 1SD @30d, M: @45d
        ws.cell(row=r, column=12, value=f"=ROUND(B{r}*D{r}*SQRT(30/365),2)")
        ws.cell(row=r, column=13, value=f"=ROUND(B{r}*D{r}*SQRT(45/365),2)")
        for j in range(9, 14):
            ws.cell(row=r, column=j).fill = OUT_FILL
            ws.cell(row=r, column=j).border = BORDER

    widths = [9, 9, 9, 9, 9, 9, 9, 13, 11, 9, 12, 11, 11]
    for j, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.freeze_panes = "A3"

    # ===================== 3) OptionScorer =====================
    ws = wb.create_sheet("OptionScorer")

    # 參數區(可調)
    ws["A1"] = "參數(可改)"
    ws["A1"].font = Font(bold=True, color="2A75BB")
    params = [
        ("最低OI", "min_oi", 500),
        ("最低量", "min_vol", 50),
        ("最大價差%", "max_spr", 0.06),
        ("買-Delta低", "buy_lo", 0.55),
        ("買-Delta高", "buy_hi", 0.70),
        ("賣-Delta低", "sell_lo", 0.15),
        ("賣-Delta高", "sell_hi", 0.30),
    ]
    for i, (label, name, val) in enumerate(params, 2):
        ws.cell(row=i, column=1, value=label).font = Font(size=9)
        c = ws.cell(row=i, column=2, value=val)
        c.fill = IN_FILL
        c.border = BORDER
    # 命名參數位置: B2..B8
    P = {name: f"$B${i}" for i, (_, name, _) in enumerate(params, 2)}

    hrow = 11
    headers = ["Ticker", "Spot", "RefIV", "HV20", "Side", "Strike", "DTE",
               "|Delta|", "OI", "Volume", "Bid", "Ask", "IV",
               "↓自動↓ Mid", "Spread%", "1SD", "距Strike",
               "流動性25", "Delta25", "到價25", "IV25", "總分", "過濾"]
    for j, h in enumerate(headers, 1):
        ws.cell(row=hrow, column=j, value=h)
    style_header(ws, hrow, len(headers))

    # 對應 demo 嘅幾個候選(同 screener 對照)
    rows = [
        # Ticker Spot RefIV HV20 Side Strike DTE |D|  OI   Vol Bid  Ask  IV
        ["AAPL", 195, 0.22, 0.234, "BUY",  190, 35, 0.68, 2983, 800, 8.46, 8.82, 0.22],
        ["AAPL", 195, 0.22, 0.234, "BUY",  200, 45, 0.59, 3059, 760, 8.05, 8.39, 0.22],
        ["NVDA", 120, 0.72, 0.544, "SELL", 111, 35, 0.32, 1768, 430, 6.39, 6.59, 0.74],
        ["NVDA", 120, 0.72, 0.544, "SELL", 108, 35, 0.28, 1303, 360, 5.39, 5.59, 0.76],
        ["KO",   62, 0.16, 0.143, "SELL",  58, 35, 0.22, 2100, 300, 0.55, 0.62, 0.17],
        ["SPY",  540, 0.12, 0.158, "BUY",  535, 45, 0.60, 5000, 1200, 12.1, 12.4, 0.12],
    ]
    d0 = hrow + 1
    for r, row in enumerate(rows, d0):
        for j, v in enumerate(row, 1):
            cell = ws.cell(row=r, column=j, value=v)
            cell.fill = IN_FILL
            cell.border = BORDER
        # 自動欄
        # N Mid
        ws.cell(row=r, column=14, value=f"=ROUND((K{r}+L{r})/2,2)")
        # O Spread%
        ws.cell(row=r, column=15, value=f"=IF(N{r}=0,1,ROUND((L{r}-K{r})/N{r},3))")
        # P 1SD = Spot*RefIV*sqrt(DTE/365)
        ws.cell(row=r, column=16, value=f"=ROUND(B{r}*C{r}*SQRT(G{r}/365),2)")
        # Q 距Strike
        ws.cell(row=r, column=17, value=f"=ABS(F{r}-B{r})")
        # R 流動性 = MIN(25,(OI/2000)*12+(Vol/500)*8+(1-MIN(spr/maxspr,1))*5)
        ws.cell(row=r, column=18,
                value=(f"=ROUND(MIN(25,(I{r}/2000)*12+(J{r}/500)*8"
                       f"+(1-MIN(O{r}/{P['max_spr']},1))*5),1)"))
        # S Delta分
        lo = f"IF(E{r}=\"BUY\",{P['buy_lo']},{P['sell_lo']})"
        hi = f"IF(E{r}=\"BUY\",{P['buy_hi']},{P['sell_hi']})"
        ws.cell(row=r, column=19,
                value=(f"=ROUND(IF(AND(H{r}>={lo},H{r}<={hi}),25,"
                       f"MAX(0,25-MIN(ABS(H{r}-({lo})),ABS(H{r}-({hi})))*100)),1)"))
        # T 到價分(buy/sell 唔同)
        ws.cell(row=r, column=20,
                value=(f"=ROUND(IF(P{r}<=0,0,"
                       f"IF(E{r}=\"BUY\","
                       f"IF(Q{r}<=P{r},25,MAX(0,25-(Q{r}-P{r})/P{r}*25)),"
                       f"IF(Q{r}>=P{r},25,IF(Q{r}>=0.5*P{r},"
                       f"12+(Q{r}-0.5*P{r})/(0.5*P{r})*13,"
                       f"MAX(0,Q{r}/(0.5*P{r})*12))))),1)"))
        # U IV regime 分: ratio=RefIV/HV20
        ws.cell(row=r, column=21,
                value=(f"=ROUND(IF(D{r}=0,0,IF(E{r}=\"SELL\","
                       f"MAX(0,MIN(25,(C{r}/D{r}-1)/0.5*25)),"
                       f"MAX(0,MIN(25,(1.2-C{r}/D{r})/0.5*25)))),1)"))
        # V 總分
        ws.cell(row=r, column=22, value=f"=ROUND(R{r}+S{r}+T{r}+U{r},1)")
        # W 過濾
        ws.cell(row=r, column=23,
                value=(f"=IF(AND(I{r}>={P['min_oi']},J{r}>={P['min_vol']},"
                       f"O{r}<={P['max_spr']}),\"✓\",\"✗\")"))
        for j in range(14, 24):
            ws.cell(row=r, column=j).fill = OUT_FILL
            ws.cell(row=r, column=j).border = BORDER

    widths = [7, 6, 6, 7, 6, 7, 5, 7, 7, 8, 6, 6, 6,
              8, 8, 7, 8, 9, 8, 8, 7, 7, 6]
    for j, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.freeze_panes = "A12"

    out = "options_screener_template.xlsx"
    wb.save(out)
    print(f"已生成 {out}")
    print("3 張表:說明 / Watchlist / OptionScorer")
    print("黃色=自己填, 綠色=公式自動計。可整份貼上 Google Sheets。")


if __name__ == "__main__":
    build()
